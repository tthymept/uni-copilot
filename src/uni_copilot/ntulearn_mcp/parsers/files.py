"""
One course file's bytes -> its text, page by page (or slide by slide).

In plain words:
  get_file (client/http.py) puts one file's raw bytes into a sink (memory).
  This file turns those bytes into text that Claude and search can use:
  - PDF: the text of each page (pypdf). Pages that are only pictures or
    scans have no text (no OCR, a plan non-goal) and are counted as empty.
  - PPTX: per slide, the text boxes, tables and speaker notes (python-pptx).
  - Jupyter notebooks (.ipynb, plain JSON): Markdown and code cells, grouped
    into "pages" that start at each heading; outputs are left out.
  - Plain text files (.txt, .csv, .md, .py, ...): split into page-sized parts.
  Anything else (old .ppt, .docx, zip, ...) gets a plain "not supported yet".

Course files are written by other people, so this file only READS them:
nothing in them is run, no macro, no other program is started. Two guards:
- A PPTX is a zip archive. A booby-trapped zip ("zip bomb") can be small to
  download but enormous unpacked, so its unpacked size is checked first.
- Damaged or password-protected files raise UnreadableFile with a plain
  message (a ValueError, which the tools pass on as it is), never a crash.

Which type a file is: the first bytes decide where they can ("%PDF-" is a
PDF; "PK" is a zip, a PPTX if it holds ppt/presentation.xml), then the file
name's extension, then the mime type NTULearn gave (the probe saw notebooks
labelled application/octet-stream, so the label alone isn't trusted).
"""

import io
import json
import logging
import re
import zipfile
from dataclasses import dataclass
from pathlib import PurePosixPath

from uni_copilot.ntulearn_mcp.models import FileType

log = logging.getLogger(__name__)

# A PPTX may unpack to at most this much (the zip-bomb guard).
MAX_UNPACKED_BYTES = 300 * 1_000_000
# Plain text and notebooks are cut into "pages" of about this many characters.
TEXT_PAGE_CHARS = 3000
TEXT_EXTENSIONS = {".txt", ".csv", ".md", ".py", ".java", ".c", ".cpp", ".r", ".m", ".sql", ".json", ".tex"}
# Types seen on your courses that we don't read (yet), for a clearer message.
KNOWN_UNSUPPORTED = {".ppt": "old PowerPoint files (.ppt)", ".pps": "old PowerPoint shows (.pps)",
                     ".docx": "Word files (.docx)", ".doc": "old Word files (.doc)", ".zip": "zip archives",
                     ".xlsx": "Excel files (.xlsx)"}

# pypdf writes warnings for every small flaw in a PDF; keep only real errors.
logging.getLogger("pypdf").setLevel(logging.ERROR)


# The file type can't be read (yet). The message names the type in plain words.
class UnsupportedFile(ValueError):
    pass


# The file is the right type but can't be read: damaged, password-protected,
# or suspicious (the zip-bomb guard).
class UnreadableFile(ValueError):
    pass


# The text of one file: one string per page/slide/part, in order. An empty
# string is a page with no text (a picture or a scan).
@dataclass(frozen=True)
class ExtractedText:
    file_type: FileType
    pages: list[str]

    # How many pages have no text at all.
    @property
    def empty_pages(self) -> int:
        return sum(1 for p in self.pages if not p.strip())


# ---------- Which type is it? ----------

# Decides the file type from the first bytes, the name and the mime type (in
# that order of trust). Raises UnsupportedFile with a plain message otherwise.
def detect_type(data: bytes, name: str = "", mime: str | None = None) -> FileType:
    ext = PurePosixPath(name.lower()).suffix
    mime = (mime or "").lower()
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:2] == b"PK" and _zip_has(data, "ppt/presentation.xml"):
        return "pptx"
    if ext == ".ipynb" or (ext in ("", ".json") and _looks_like_notebook(data)):
        return "ipynb"
    if ext in TEXT_EXTENSIONS or mime.startswith("text/"):
        return "text"
    what = KNOWN_UNSUPPORTED.get(ext) or (f"'{ext}' files" if ext else f"this type ({mime or 'unknown'})")
    raise UnsupportedFile(f"Reading {what} isn't supported yet: only PDF, PPTX, notebooks and text files.")


# True if `data` is a zip archive containing a file named `member`.
def _zip_has(data: bytes, member: str) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            return member in z.namelist()
    except zipfile.BadZipFile:
        return False


# True if `data` is JSON shaped like a Jupyter notebook (a "cells" list).
def _looks_like_notebook(data: bytes) -> bool:
    try:
        doc = json.loads(data.decode("utf-8-sig"))
    except (UnicodeDecodeError, ValueError):
        return False
    return isinstance(doc, dict) and isinstance(doc.get("cells"), list)


# ---------- Reading each type ----------

# The entry point: bytes -> ExtractedText. `name` and `mime` help find the type.
def extract_text(data: bytes, name: str = "", mime: str | None = None) -> ExtractedText:
    file_type = detect_type(data, name, mime)
    reader = {"pdf": _pdf_pages, "pptx": _pptx_pages, "ipynb": _notebook_pages, "text": _text_pages}[file_type]
    pages = [_tidy(p) for p in reader(data)]
    log.debug("extracted %s: %d pages, %d empty", file_type, len(pages), sum(1 for p in pages if not p))
    return ExtractedText(file_type=file_type, pages=pages)


# PDF: one string per page, using pypdf's text extraction. A password-
# protected PDF is tried with an empty password (many "protected" PDFs only
# restrict printing); if that fails, or the file is damaged, UnreadableFile.
def _pdf_pages(data: bytes) -> list[str]:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise UnreadableFile("This PDF is password-protected, so its text can't be read.")
        pages = []
        for page in reader.pages:
            try:
                pages.append(page.extract_text() or "")
            except (PyPdfError, ValueError, KeyError, TypeError):
                pages.append("")   # one bad page shouldn't lose the rest
        return pages
    except UnreadableFile:
        raise
    except (PyPdfError, ValueError, KeyError, TypeError, OSError) as e:
        raise UnreadableFile(f"This PDF couldn't be read (it may be damaged): {type(e).__name__}.") from None


# PPTX: one string per slide: every text box and table (also inside
# grouped shapes), then the speaker notes under "Notes:". Checks the zip's
# unpacked size first (zip-bomb guard).
def _pptx_pages(data: bytes) -> list[str]:
    from pptx import Presentation

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            unpacked = sum(info.file_size for info in z.infolist())
    except zipfile.BadZipFile:
        raise UnreadableFile("This PPTX couldn't be read (it may be damaged).") from None
    if unpacked > MAX_UNPACKED_BYTES:
        raise UnreadableFile(f"This PPTX unpacks to over {MAX_UNPACKED_BYTES // 1_000_000} MB, so it wasn't read.")
    try:
        deck = Presentation(io.BytesIO(data))
    except Exception as e:   # python-pptx raises many kinds of errors on a bad file
        raise UnreadableFile(f"This PPTX couldn't be read (it may be damaged): {type(e).__name__}.") from None
    slides = []
    for slide in deck.slides:
        parts = [t for shape in slide.shapes for t in _shape_texts(shape)]
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append(f"Notes: {notes}")
        except Exception:   # odd notes shouldn't lose the slide
            log.debug("pptx: notes of one slide couldn't be read")
        slides.append("\n".join(parts))
    return slides


# The texts inside one slide shape: a text box's paragraphs, a table's rows
# as "a | b | c", and the same for every shape inside a group.
# Real decks contain shapes python-pptx doesn't know (found 30 Oct: asking
# one for its `shape_type` raised NotImplementedError). So groups are found
# with isinstance (no question asked of the shape), and each shape is read
# inside try: a strange shape loses only its own text, never the slide.
def _shape_texts(shape) -> list[str]:
    from pptx.shapes.group import GroupShape

    try:
        if isinstance(shape, GroupShape):
            return [t for inner in shape.shapes for t in _shape_texts(inner)]
        texts = []
        if getattr(shape, "has_text_frame", False):
            texts += [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    texts.append(" | ".join(cells))
        return texts
    except Exception:   # python-pptx raises many kinds of errors on unusual shapes
        log.debug("pptx: one shape couldn't be read (%s)", type(shape).__name__)
        return []


# Notebook: Markdown cells as they are, code cells in ``` fences, outputs left
# out (they can be huge and are rarely the point). A new "page" starts at
# each Markdown heading, or when a page passes TEXT_PAGE_CHARS.
def _notebook_pages(data: bytes) -> list[str]:
    try:
        doc = json.loads(data.decode("utf-8-sig"))
        cells = doc["cells"]
    except (UnicodeDecodeError, ValueError, KeyError, TypeError):
        raise UnreadableFile("This notebook couldn't be read (it isn't valid notebook JSON).") from None
    pages: list[str] = []
    current: list[str] = []
    for cell in cells if isinstance(cells, list) else []:
        if not isinstance(cell, dict):
            continue
        source = cell.get("source", "")
        text = "".join(source) if isinstance(source, list) else str(source)
        if not text.strip():
            continue
        if cell.get("cell_type") == "code":
            text = f"```\n{text.rstrip()}\n```"
        starts_section = cell.get("cell_type") == "markdown" and text.lstrip().startswith("#")
        if current and (starts_section or sum(len(t) for t in current) > TEXT_PAGE_CHARS):
            pages.append("\n\n".join(current))
            current = []
        current.append(text.strip())
    if current:
        pages.append("\n\n".join(current))
    return pages or [""]


# Plain text: decoded as UTF-8 (unreadable bytes become a replacement mark,
# never an error) and cut at line breaks into parts of about TEXT_PAGE_CHARS.
def _text_pages(data: bytes) -> list[str]:
    text = data.decode("utf-8-sig", errors="replace")
    pages: list[str] = []
    current = ""
    for line in text.splitlines(keepends=True):
        if current and len(current) + len(line) > TEXT_PAGE_CHARS:
            pages.append(current)
            current = ""
        current += line
    if current or not pages:
        pages.append(current)
    return pages


# Cleans one page's text: no NUL characters, no trailing spaces on lines,
# at most one empty line in a row.
def _tidy(text: str) -> str:
    text = text.replace("\x00", "")
    lines = [line.rstrip() for line in text.splitlines()]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


# ---------- Choosing pages for one read_file_text call ----------

# Turns the tool's `pages` argument ("5", "3-10", "21-", None) into the page
# numbers to return, keeping within max_pages and max_chars. Returns
# (page numbers, cut_short, next_pages) where next_pages is what to ask for
# next ("21-40"), or None when nothing is left. Raises ValueError for a bad
# range, with a plain message.
def choose_pages(spec: str | None, page_texts: list[str], max_pages: int, max_chars: int) -> tuple[list[int], bool, str | None]:
    total = len(page_texts)
    first, last = _parse_range(spec, total)
    chosen: list[int] = []
    chars = 0
    for number in range(first, last + 1):
        size = len(page_texts[number - 1])
        if chosen and (len(chosen) >= max_pages or chars + size > max_chars):
            break
        chosen.append(number)
        chars += size
    cut_short = bool(chosen) and chosen[-1] < last
    next_pages = None
    if chosen and chosen[-1] < total:
        start = chosen[-1] + 1
        next_pages = f"{start}-{min(total, start + max_pages - 1)}"
    return chosen, cut_short, next_pages


# "5" -> (5, 5); "3-10" -> (3, 10); "21-" -> (21, total); None -> (1, total).
# Numbers are clamped to the file's pages; a range that starts after the end
# is an error ("This file has 12 pages.").
def _parse_range(spec: str | None, total: int) -> tuple[int, int]:
    if total == 0:
        raise ValueError("This file has no pages.")
    if not spec or not spec.strip():
        return 1, total
    m = re.fullmatch(r"\s*(\d+)\s*(?:(-)\s*(\d*)\s*)?", spec)
    if not m:
        raise ValueError(f"pages must look like '5', '3-10' or '21-', not {spec!r}.")
    first = int(m.group(1))
    last = int(m.group(3)) if m.group(3) else (total if m.group(2) else first)
    if first < 1 or first > total or last < first:
        raise ValueError(f"No such pages: {spec!r}. This file has {total} page{'s' if total != 1 else ''}.")
    return first, min(last, total)


# ---------- Self-test ----------

# A tiny but real PDF, written by hand: one page per text in `texts` ("" = an
# empty page), using the built-in Helvetica font. Only for the self-test and
# tests; we never write PDFs otherwise.
def _tiny_pdf(texts: list[str]) -> bytes:
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", None,
               "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in texts:
        stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET" if text else ""
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        content_ref = len(objects)
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                       f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_ref} 0 R >>")
        kids.append(f"{len(objects)} 0 R")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(f"{n} 0 obj\n{body}\nendobj\n".encode())
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


# A tiny PPTX built with python-pptx: a title + bullet slide with notes, and
# a slide with a 2x2 table. Only for the self-test and tests.
def _tiny_pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    s1 = deck.slides.add_slide(deck.slide_layouts[1])
    s1.shapes.title.text = "Week 10: Shortest paths"
    s1.placeholders[1].text = "Dijkstra's algorithm"
    s1.notes_slide.notes_text_frame.text = "Mention the quiz on Friday."
    s2 = deck.slides.add_slide(deck.slide_layouts[5])
    s2.shapes.title.text = "Complexity"
    table = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
    for r, row in enumerate([["Algorithm", "Time"], ["Dijkstra", "O(E log V)"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    group = s2.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(4), Inches(4), Inches(1)).text_frame.text = "Text inside a group"
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


# Two "does this text look jumbled?" numbers, so text quality can be judged
# without printing any of it: the share of one-letter words (broken spacing
# like "D i j k s t r a" makes it high; normal English is a few percent) and
# the share of words over 30 characters (missing spaces glue words together).
def _jumble_signs(text: str) -> tuple[float, float]:
    words = text.split()
    if not words:
        return 0.0, 0.0
    single = sum(1 for w in words if len(w) == 1 and w.isalpha())
    long = sum(1 for w in words if len(w) > 30)
    return single / len(words), long / len(words)


# Live check (--real): up to 3 real files of each kind from your courses (a
# PDF / PPTX file item, a PDF / PPTX / notebook embedded in a page), fetched
# into memory with get_file, read, and dropped; stops at the first that works
# per kind. Prints counts and the jumble signs only: never names, links or
# text. If an embedded file fails, shows why without the link: its masked
# shape and each hop's host and status (the probe's safe helpers).
# GET only, ~40-60 requests.
def _real_check() -> None:
    import time

    from uni_copilot.ntulearn_mcp.client.courses import CoursesClient
    from uni_copilot.ntulearn_mcp.client.http import NtuLearnError, NtuLearnHttp
    from uni_copilot.ntulearn_mcp.config import READ_MAX_BYTES, load_settings
    from uni_copilot.ntulearn_mcp.parsers.content import embedded_files, embedded_links, handler
    from uni_copilot.ntulearn_mcp.probe_content import follow_download, masked_shape

    http = NtuLearnHttp()
    sites = [s.id for c in CoursesClient(http, load_settings()).list_courses().courses for s in c.sites]
    wanted = ["PDF file item", "PPTX file item", "embedded PDF", "embedded PPTX", "embedded notebook"]
    found: dict[str, list[tuple[str, str, str | None]]] = {k: [] for k in wanted}   # label -> [(path or link, name, mime)]
    no_link: dict[str, int] = {}   # extension -> embedded files with no usable link
    for site in sites:
        tree = http.get_all(f"/learn/api/public/v1/courses/{site}/contents", {"recursive": "true", "limit": 200})
        for raw in tree:
            info = (raw.get("contentHandler") or {}).get("file") or {}
            name = str(info.get("fileName") or "")
            kind = {".pdf": "PDF file item", ".pptx": "PPTX file item"}.get(PurePosixPath(name.lower()).suffix)
            if handler(raw) == "resource/x-bb-file" and kind and not found[kind]:
                atts = http.get_all(f"/learn/api/public/v1/courses/{site}/contents/{raw['id']}/attachments")
                if atts:
                    path = f"/learn/api/public/v1/courses/{site}/contents/{raw['id']}/attachments/{atts[0]['id']}/download"
                    found[kind].append((path, name, atts[0].get("mimeType")))
        # Every document with a body: Ultra pages' hidden bodies and the other
        # documents (the probe's 85, which have their own body), not only the former.
        for page_doc in (raw for raw in tree if handler(raw) == "resource/x-bb-document" and raw.get("body")):
            links = embedded_links(page_doc.get("body"))
            for f in embedded_files(page_doc.get("body")):
                ext = PurePosixPath(f.name.lower()).suffix
                if not f.has_link:
                    no_link[ext or "(none)"] = no_link.get(ext or "(none)", 0) + 1
                kind = {".pdf": "embedded PDF", ".pptx": "embedded PPTX", ".ipynb": "embedded notebook"}.get(ext)
                if kind and len(found[kind]) < 3 and f.xid in links:
                    found[kind].append((links[f.xid], f.name, f.mime_type))

    print("Real files (GET only, read in memory, nothing saved; counts only):")
    print(f"   embedded files with no usable link, by extension: {no_link or 'none'}")
    for label in wanted:
        if not found[label]:
            print(f"   {label}: none found on your courses")
        for n, (start, name, mime) in enumerate(found[label], 1):
            sink = io.BytesIO()
            t0 = time.monotonic()
            try:
                info = http.get_file(start, READ_MAX_BYTES, sink)
                result = extract_text(sink.getvalue(), name, mime)
            except (ValueError, NtuLearnError) as e:   # UnsupportedFile, UnreadableFile, FileTooLarge, HTTP errors
                print(f"   {label} #{n}: {e}")
                if label.startswith("embedded"):
                    print(f"      link shape: {masked_shape(start)}")
                    follow_download(http, start)
                continue
            text = "\n".join(result.pages)
            single, long = _jumble_signs(text)
            print(f"   {label} #{n}: {info.size / 1_000_000:.2f} MB -> {result.file_type}, {len(result.pages)} pages, "
                  f"{result.empty_pages} empty, {len(text):,} characters; one-letter words {single:.1%}, "
                  f"words over 30 characters {long:.1%}; {time.monotonic() - t0:.1f} s")
            break


# Self-test: uv run python -m uni_copilot.ntulearn_mcp.parsers.files [--real]
# Without --real: offline, tiny made-up files of every type read back.
def _main() -> None:
    import argparse

    ap = argparse.ArgumentParser(description="Self-test of the file text reader.")
    ap.add_argument("--real", action="store_true", help="live: read one real file of each kind (GET only, counts only)")
    if ap.parse_args().real:
        logging.basicConfig(level=logging.WARNING)
        _real_check()
        return
    pdf = extract_text(_tiny_pdf(["Breadth-first search", "", "Dijkstra and A-star"]), "made-up.pdf")
    print(f"1) PDF: {len(pdf.pages)} pages, {pdf.empty_pages} empty; page 1 {pdf.pages[0]!r}, page 3 {pdf.pages[2]!r}")

    deck = extract_text(_tiny_pptx(), "made-up.pptx", "application/octet-stream")
    print(f"2) PPTX (labelled octet-stream, found by its bytes): {len(deck.pages)} slides")
    for n, slide in enumerate(deck.pages, 1):
        print(f"   slide {n}: {slide!r}")

    # A shape python-pptx doesn't know raises on any question (seen in a real deck, 30 Oct).
    class _UnknownShape:
        @property
        def has_text_frame(self):
            raise NotImplementedError("Shape instance of unrecognized shape type")
    print(f"   an unknown shape -> {_shape_texts(_UnknownShape())} (skipped, no crash)")

    notebook = json.dumps({"cells": [
        {"cell_type": "markdown", "source": ["# Lab 4\n", "Load the data."]},
        {"cell_type": "code", "source": "import pandas as pd", "outputs": [{"text": "huge output"}]},
        {"cell_type": "markdown", "source": "## Part 2\nPlot it."},
    ]}).encode()
    nb = extract_text(notebook, "lab4.ipynb", "application/octet-stream")
    print(f"3) Notebook: {len(nb.pages)} parts (split at headings), outputs left out: {'huge output' not in ''.join(nb.pages)}")

    text = extract_text(("line of a csv,1,2\n" * 400).encode(), "data.csv", "text/csv")
    print(f"4) Text: {len(text.pages)} parts of about {TEXT_PAGE_CHARS} characters")

    for name, data in (("slides.ppt", b"\xd0\xcf\x11\xe0 old office file"), ("archive.zip", b"PK\x03\x04 not a pptx")):
        try:
            extract_text(data, name)
        except UnsupportedFile as e:
            print(f"5) {name}: {e}")
    try:
        extract_text(b"%PDF-1.4 this is not really a PDF", "broken.pdf")
    except UnreadableFile as e:
        print(f"6) broken.pdf: {e}")

    pages = [f"page {n} " + "x" * 2000 for n in range(1, 31)]
    print("7) choose_pages:", choose_pages(None, pages, 20, 30_000)[1:], choose_pages("21-", pages, 20, 30_000),
          choose_pages("5", pages, 20, 30_000))
    try:
        choose_pages("40-45", pages, 20, 30_000)
    except ValueError as e:
        print("   '40-45' ->", e)


if __name__ == "__main__":
    _main()
